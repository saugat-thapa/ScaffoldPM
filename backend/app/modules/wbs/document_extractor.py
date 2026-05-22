import io
import re
from typing import Iterable

import pdfplumber
import fitz
from docx import Document
from fastapi import HTTPException


_PAGE_NUMBER_RE = re.compile(r"^\s*(page\s*)?\d+(\s*of\s*\d+)?\s*$", re.IGNORECASE)
_LIST_ITEM_RE = re.compile(r"^(\d+[\.)]|[-*\u2022])\s+")


def extract_document_text(filename: str | None, content_type: str | None, data: bytes) -> str:
    ext = (filename or "").lower().split(".")[-1] if filename and "." in filename else ""
    content_type = (content_type or "").lower()

    if ext == "pdf" or content_type == "application/pdf":
        pages = _extract_pdf_pages(data)
    elif ext in {"docx", "doc"} or content_type in {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    }:
        pages = _extract_docx_pages(data)
    elif ext == "txt" or content_type.startswith("text/"):
        pages = [_extract_txt(data)]
    else:
        raise HTTPException(status_code=400, detail="Unsupported file type")

    return _clean_pages(pages)


def _extract_pdf_pages(data: bytes) -> list[str]:
    pages: list[str] = []
    fitz_doc: fitz.Document | None = None

    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for idx, page in enumerate(pdf.pages):
                text = page.extract_text() or ""
                if not text.strip():
                    if fitz_doc is None:
                        fitz_doc = fitz.open(stream=data, filetype="pdf")
                    if idx < fitz_doc.page_count:
                        text = fitz_doc.load_page(idx).get_text("text") or ""
                pages.append(text)
    except Exception:
        if fitz_doc is None:
            fitz_doc = fitz.open(stream=data, filetype="pdf")
        for idx in range(fitz_doc.page_count):
            pages.append(fitz_doc.load_page(idx).get_text("text") or "")

    if fitz_doc is not None:
        fitz_doc.close()

    return pages


def _extract_docx_pages(data: bytes) -> list[str]:
    doc = Document(io.BytesIO(data))
    lines: list[str] = []

    for para in doc.paragraphs:
        if para.text and para.text.strip():
            lines.append(para.text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    return ["\n".join(lines)]


def _extract_txt(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _clean_pages(pages: Iterable[str]) -> str:
    normalized_pages: list[list[str]] = []
    header_candidates: list[str] = []
    footer_candidates: list[str] = []

    for raw in pages:
        lines = [_normalize_line(line) for line in raw.splitlines()]
        lines = [line for line in lines if line]
        normalized_pages.append(lines)
        if lines:
            header_candidates.append(lines[0])
            footer_candidates.append(lines[-1])

    repeated_headers = _repeated_lines(header_candidates)
    repeated_footers = _repeated_lines(footer_candidates)
    repeated_noise = repeated_headers | repeated_footers

    cleaned_lines: list[str] = []
    for lines in normalized_pages:
        for line in lines:
            if line in repeated_noise:
                continue
            if _PAGE_NUMBER_RE.match(line):
                continue
            cleaned_lines.append(line)
        cleaned_lines.append("")

    return _reflow_lines(cleaned_lines).strip()


def _normalize_line(line: str) -> str:
    line = line.replace("\t", " ").strip()
    line = re.sub(r"\s+", " ", line)
    return line


def _repeated_lines(lines: Iterable[str]) -> set[str]:
    counts: dict[str, int] = {}
    for line in lines:
        if not line:
            continue
        counts[line] = counts.get(line, 0) + 1
    return {line for line, count in counts.items() if count > 1}


def _is_heading(line: str) -> bool:
    if len(line) > 80:
        return False
    alpha = sum(1 for ch in line if ch.isalpha())
    if alpha < 3:
        return False
    return line.isupper() or line.endswith(":")


def _is_list_item(line: str) -> bool:
    return bool(_LIST_ITEM_RE.match(line))


def _reflow_lines(lines: Iterable[str]) -> str:
    paragraphs: list[str] = []
    buffer = ""

    for line in lines:
        if not line:
            if buffer:
                paragraphs.append(buffer)
                buffer = ""
            continue

        if _is_heading(line) or _is_list_item(line):
            if buffer:
                paragraphs.append(buffer)
                buffer = ""
            paragraphs.append(line)
            continue

        if not buffer:
            buffer = line
            continue

        if buffer.endswith("-") and not buffer.endswith("--"):
            buffer = buffer[:-1] + line
        else:
            buffer = f"{buffer} {line}"

    if buffer:
        paragraphs.append(buffer)

    return "\n\n".join(paragraphs)
