from __future__ import annotations

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


REPORT = "ScaffoldPM_Final_Year_Project_Report_Chapters_1_to_4.docx"
FIGURES = [
    "Figure 1: Use Case Diagram for ScaffoldPM",
    "Figure 2: Gantt Chart for ScaffoldPM Development",
    "Figure 3: Class Diagram for ScaffoldPM",
    "Figure 4: Sequence Diagram for Project Plan Generation",
    "Figure 5: Activity Diagram for AI-Assisted Project Planning Pipeline",
    "Figure 6: State Diagram for Project Generation Request Lifecycle",
    "Figure 7: System Flowchart for ScaffoldPM",
    "Figure 8: Form and Report Design for ScaffoldPM",
    "Figure 9: Interface and Dialogue Design for ScaffoldPM",
]


def set_font(paragraph) -> None:
    for run in paragraph.runs:
        run.font.name = "Calibri"
        if run._element.rPr is not None:
            run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        run.font.size = Pt(11)


def remove_paragraphs(paragraphs) -> None:
    for paragraph in paragraphs:
        paragraph._element.getparent().remove(paragraph._element)


def main() -> None:
    doc = Document(REPORT)
    paragraphs = list(doc.paragraphs)

    toc_start = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "TABLE OF CONTENTS")
    actual_abbrev = next(i for i, p in enumerate(paragraphs) if i > toc_start and p.style.name == "Heading 1" and p.text.strip() == "LIST OF ABBREVIATIONS")
    remove_paragraphs([p for p in paragraphs[toc_start:actual_abbrev] if p.text.strip().startswith("Figure ")])

    paragraphs = list(doc.paragraphs)
    actual_figs = next(i for i, p in enumerate(paragraphs) if p.style.name == "Heading 1" and p.text.strip() == "LIST OF FIGURES")
    actual_tables = next(i for i, p in enumerate(paragraphs) if i > actual_figs and p.style.name == "Heading 1" and p.text.strip() == "LIST OF TABLES")
    remove_paragraphs(paragraphs[actual_figs + 1:actual_tables])

    paragraphs = list(doc.paragraphs)
    actual_tables_para = next(p for p in paragraphs if p.style.name == "Heading 1" and p.text.strip() == "LIST OF TABLES")
    for item in FIGURES:
        p = actual_tables_para.insert_paragraph_before(item)
        set_font(p)

    doc.save(REPORT)


if __name__ == "__main__":
    main()
