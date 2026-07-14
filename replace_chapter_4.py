from __future__ import annotations

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPORT = "ScaffoldPM_Final_Year_Project_Report_Chapters_1_to_4.docx"
OUT_REPORT = "ScaffoldPM_Final_Year_Project_Report_Chapters_1_to_4_Updated_Chapter_4.docx"


FIGURES = [
    "Figure 1: Use Case Diagram for ScaffoldPM",
    "Figure 2: Gantt Chart for ScaffoldPM Development",
    "Figure 3: Class Diagram for ScaffoldPM",
    "Figure 4: Sequence Diagram for Project Plan Generation",
    "Figure 5: Activity Diagram for AI-Assisted Project Planning Pipeline",
    "Figure 6: State Diagram for Project Generation Request Lifecycle",
    "Figure 7: Component Diagram for ScaffoldPM",
    "Figure 8: Deployment Diagram for ScaffoldPM",
    "Figure 9: Frontend Interface Design for ScaffoldPM",
]


def set_font(run, size=11, bold=False, italic=False, color=None, name="Calibri") -> None:
    run.font.name = name
    if run._element.rPr is not None:
        run._element.rPr.rFonts.set(qn("w:ascii"), name)
        run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def format_table(table, widths: list[float]) -> None:
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for i, width in enumerate(widths):
            cell = row.cells[i]
            cell.width = Inches(width)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def add_before(anchor, text: str, style: str | None = None, justify=False, italic=False):
    p = anchor.insert_paragraph_before(text)
    if style:
        p.style = style
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs:
        set_font(r, italic=italic)
    return p


def add_bullet(anchor, text: str) -> None:
    p = add_before(anchor, text, style="List Bullet")
    for r in p.runs:
        set_font(r)


def add_number(anchor, text: str) -> None:
    p = add_before(anchor, text, style="List Number")
    for r in p.runs:
        set_font(r)


def add_table_before(anchor, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    p = anchor.insert_paragraph_before("")
    table = p._parent.add_table(rows=1, cols=len(headers), width=Inches(6.5))
    p._p.addnext(table._tbl)
    format_table(table, widths)
    set_repeat_table_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "E8EEF5")
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(header)
        set_font(run, size=10, bold=True, color="0B2545")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            para = cells[i].paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = para.add_run(value)
            set_font(run, size=10)
    anchor.insert_paragraph_before("")


def add_placeholder(anchor, title: str, description: str) -> None:
    p = add_before(anchor, title)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        set_font(r, size=10, bold=True, color="1F4D78")
    desc = add_before(anchor, f"[Insert {title.lower()} here. Suggested image description: {description}]", justify=True, italic=True)
    for r in desc.runs:
        set_font(r, size=10, italic=True, color="555555")


def clear_runs(paragraph, text: str) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    run = paragraph.add_run(text)
    set_font(run)


def remove_between(doc: Document, start_text: str, end_text: str, after_text: str | None = None):
    paragraphs = list(doc.paragraphs)
    after_i = -1
    if after_text is not None:
        after_i = next(
            i for i, p in enumerate(paragraphs)
            if p.style.name == "Heading 1" and p.text.strip() == after_text
        )
    start_i = next(
        i for i, p in enumerate(paragraphs)
        if i > after_i and p.style.name == "Heading 1" and p.text.strip() == start_text
    )
    end_i = next(
        i for i, p in enumerate(paragraphs)
        if i > start_i and p.style.name == "Heading 1" and p.text.strip() == end_text
    )
    start_p = paragraphs[start_i]
    end_p = paragraphs[end_i]
    body = doc._body._element
    children = list(body)
    start_child = children.index(start_p._element)
    end_child = children.index(end_p._element)
    for element in children[start_child:end_child]:
        body.remove(element)
    return end_p


def update_figures(doc: Document) -> None:
    paragraphs = list(doc.paragraphs)
    actual_figs = next(i for i, p in enumerate(paragraphs) if p.style.name == "Heading 1" and p.text.strip() == "LIST OF FIGURES")
    actual_tables = next(i for i, p in enumerate(paragraphs) if i > actual_figs and p.style.name == "Heading 1" and p.text.strip() == "LIST OF TABLES")
    for p in paragraphs[actual_figs + 1:actual_tables]:
        p._element.getparent().remove(p._element)

    paragraphs = list(doc.paragraphs)
    tables_heading = next(p for p in paragraphs if p.style.name == "Heading 1" and p.text.strip() == "LIST OF TABLES")
    for item in FIGURES:
        p = tables_heading.insert_paragraph_before(item)
        for r in p.runs:
            set_font(r)


def add_chapter_4(anchor) -> None:
    add_before(anchor, "CHAPTER 4 - SYSTEM DESIGN", style="Heading 1")
    add_before(
        anchor,
        "This chapter presents the detailed design of ScaffoldPM, refining the analysis models into implementation-ready specifications for an AI-driven project planning backend. The design addresses the complete workflow from scope document upload through document extraction, LLM-based WBS generation, semantic task assignment, duration estimation, dependency-aware scheduling, critical path analysis, database persistence, and JSON response delivery.",
        justify=True,
    )

    add_before(anchor, "Architectural Design", style="Heading 2")
    add_before(
        anchor,
        "ScaffoldPM follows a modular service-layer architecture with a clear separation between the API layer, processing modules, and database layer. The backend is implemented using FastAPI and exposes REST endpoints for project generation, project retrieval, task status update, and task reassignment.",
        justify=True,
    )
    add_bullet(anchor, "API Layer: Handles HTTP routing, multipart file upload, payload validation, error handling, and JSON response construction. The project.py router defines the main /api/project endpoints, while health.py provides backend health and extraction testing endpoints.")
    add_bullet(anchor, "Service Layer: Contains the main processing modules. The document extractor reads PDF, DOC/DOCX, and TXT files; the WBS generator creates project structure using an LLM; the matcher assigns tasks using semantic embeddings; the predictor estimates duration; the scheduler assigns dates; and the CPM analyzer identifies critical tasks.")
    add_bullet(anchor, "Data Layer: Contains Pydantic schemas and SQLAlchemy models. Pydantic models validate and serialize project objects, while SQLAlchemy models define database tables for projects, phases, activities, tasks, team members, and project-member relationships.")
    add_bullet(anchor, "Infrastructure Layer: Provides configuration, OpenAI API connectivity, asynchronous SQLite access, model loading for duration prediction, SentenceTransformer embedding support, document parsing libraries, and NetworkX graph processing.")

    add_before(anchor, "Component Diagram", style="Heading 3")
    add_before(anchor, "The component diagram describes the major software components of ScaffoldPM and how they interact during project generation.", justify=True)
    add_bullet(anchor, "Client Tier: A future web frontend or API testing client that uploads the scope document and sends project metadata to the backend.")
    add_bullet(anchor, "API Tier: FastAPI application server running the ScaffoldPM backend. It exposes routes under /api and /api/project and provides OpenAPI documentation through FastAPI.")
    add_bullet(anchor, "Processing Tier: DocumentExtractor, WBSGenerator, TaskMatcher, DurationPredictor, Scheduler, and CPMAnalyzer modules.")
    add_bullet(anchor, "External AI Service: OpenAI LLM service used by the WBS generation module through LangChain.")
    add_bullet(anchor, "Persistence Tier: SQLAlchemy asynchronous database access layer and SQLite database file.")
    add_bullet(anchor, "Interfaces: ProjectRouter communicates with DocumentExtractor, WBSGenerator, TaskMatcher, DurationPredictor, Scheduler, CPMAnalyzer, and ProjectRepository. WBSGenerator communicates with OpenAI API. ProjectRepository communicates with SQLite.")
    add_placeholder(
        anchor,
        "Figure 7: Component Diagram for ScaffoldPM",
        "Create a UML component diagram with Client/API Client, FastAPI App, ProjectRouter, DocumentExtractor, WBSGenerator, OpenAI LLM Service, TaskMatcher, DurationPredictor, Scheduler, CPMAnalyzer, ProjectRepository, SQLAlchemy ORM, and SQLite Database. Show REST call from client to API, service calls between router and modules, external call to OpenAI, and persistence call to SQLite.",
    )

    add_before(anchor, "Deployment Diagram", style="Heading 3")
    add_before(anchor, "The deployment design explains how ScaffoldPM can be hosted in a practical environment. The current implementation can run locally for academic demonstration, but the design also supports future deployment to a server environment.", justify=True)
    add_bullet(anchor, "Node 1 - Client Device: Laptop or desktop browser, Swagger UI, Postman, or a future frontend application. Communicates with the backend through HTTP requests.")
    add_bullet(anchor, "Node 2 - API Server: Python runtime running FastAPI with Uvicorn ASGI server. Recommended configuration is Python 3.12 or compatible version, 4 vCPU, 8 GB RAM, and SSD storage.")
    add_bullet(anchor, "Node 3 - AI/ML Runtime: Executes SentenceTransformer embeddings, XGBoost duration prediction, and document processing libraries. The LLM call is sent to the external OpenAI API over the internet.")
    add_bullet(anchor, "Node 4 - Database Storage: SQLite database file scaffoldpm.db for local persistence. In production, this can be replaced with PostgreSQL or another managed relational database.")
    add_bullet(anchor, "Communication: Client communicates with FastAPI over HTTP/HTTPS; FastAPI communicates with OpenAI API over HTTPS; FastAPI accesses the SQLite database through SQLAlchemy and aiosqlite.")
    add_placeholder(
        anchor,
        "Figure 8: Deployment Diagram for ScaffoldPM",
        "Create a UML deployment diagram with Client Device, API Server running FastAPI/Uvicorn, AI/ML Runtime modules, External OpenAI API, and SQLite Database Storage. Show HTTP/HTTPS communication, local module execution, and database access.",
    )

    add_before(anchor, "Data Model Design", style="Heading 2")
    add_before(anchor, "The data model design stores the generated project plan in a relational structure. A project contains phases, phases contain activities, and activities contain tasks. Team members are linked to projects through a junction table and can be assigned to tasks. The schema is suitable for storing WBS hierarchy, schedule information, assignment information, and critical path results.", justify=True)
    add_table_before(anchor, ["Table", "Column", "Type", "Constraints / Description"], [
        ["projects", "id, name", "String", "Primary key; project name is required"],
        ["projects", "description, summary", "Text", "Optional project and AI summary fields"],
        ["projects", "start_date, end_date, created_at", "Date/DateTime", "Stores project timeline and creation date"],
        ["projects", "total_working_days, total_duration_days", "Integer", "Stores planned and computed working duration"],
        ["team_members", "id, name, role", "String", "Primary key; name and role are required"],
        ["team_members", "skills, experience", "Text/String", "Skills stored as comma-separated text; experience defaults to mid"],
        ["project_members", "project_id, team_member_id", "String", "Composite primary key; junction table"],
        ["phases", "id, project_id, title, order", "String/Integer", "Phase belongs to project and has display order"],
        ["activities", "id, phase_id, title", "String", "Activity belongs to phase"],
        ["tasks", "id, activity_id, assigned_to, title", "String", "Task belongs to activity and may be assigned to a team member"],
        ["tasks", "task_type, status", "Enum", "Task category and Kanban status"],
        ["tasks", "estimated_hours, start_date, end_date", "Float/Date", "Effort estimate and scheduled dates"],
        ["tasks", "dependencies, skills_required", "Text", "Comma-separated task IDs and required skills"],
        ["tasks", "is_critical, slack", "Boolean/Float", "Critical path result fields"],
    ], [1.25, 2.0, 1.25, 1.8])

    add_before(anchor, "Normalization", style="Heading 3")
    add_number(anchor, "First Normal Form (1NF): Main entities are separated into projects, phases, activities, tasks, and team members. Most attributes are atomic; however, skills and dependencies are currently stored as comma-separated text for implementation simplicity.")
    add_number(anchor, "Second Normal Form (2NF): Non-key attributes depend on the full primary key. Junction table project_members uses project_id and team_member_id to represent many-to-many membership without repeating team data inside project rows.")
    add_number(anchor, "Third Normal Form (3NF): Project details, phase details, activity details, task details, and team member details are stored in separate tables to reduce transitive dependency. Future versions can further normalize skills, task dependencies, and required skills into separate relation tables.")

    add_before(anchor, "API Interface Design", style="Heading 2")
    add_before(anchor, "The ScaffoldPM backend exposes RESTful APIs using FastAPI. The main API accepts multipart form data containing the uploaded document and JSON payload, then returns a structured project plan.", justify=True)
    add_table_before(anchor, ["Endpoint", "Method", "Input", "Output / Purpose"], [
        ["/api/health", "GET", "None", "Returns backend status, application name, and debug mode"],
        ["/api/health/extract", "POST", "Uploaded file", "Returns extracted document text for testing"],
        ["/api/project/generate-wbs", "POST", "file UploadFile; payload JSON string", "Generates WBS, assignments, schedule, total duration, and critical path"],
        ["/api/project/", "GET", "None", "Lists saved project summaries"],
        ["/api/project/{project_id}", "GET", "project_id path parameter", "Returns full saved project with phases, activities, tasks, and team"],
        ["/api/project/task/status", "PATCH", "task_id, new_status", "Updates task status to todo, in_progress, or done"],
        ["/api/project/task/assign", "PATCH", "task_id, member_id", "Reassigns a task to another team member"],
    ], [1.9, 0.75, 1.8, 1.85])

    add_table_before(anchor, ["Field", "Type", "Required", "Description"], [
        ["project_name", "string", "Yes", "Name of the project to be planned"],
        ["start_date", "date", "Yes", "Project start date"],
        ["end_date", "date", "Yes", "Project end date"],
        ["team_members", "list", "Yes", "List of team members with id, name, role, skills, and experience"],
        ["file", "UploadFile", "Yes", "Scope document in PDF, DOC, DOCX, or TXT format"],
    ], [1.5, 1.0, 1.0, 2.8])

    add_before(anchor, "Frontend Interface Design", style="Heading 2")
    add_before(anchor, "The current repository contains the backend implementation and an empty frontend directory. The frontend interface design is therefore proposed for future integration. It should consume the REST APIs provided by ScaffoldPM and present the generated plan in a user-friendly project management interface.", justify=True)
    add_bullet(anchor, "Navigation: Fixed top navigation with Project List, Generate Project, Team Members, WBS View, Schedule, Critical Path, and Settings links.")
    add_bullet(anchor, "Generate Project Page: Form fields for project name, start date, end date, scope document upload, and dynamic team member entry.")
    add_bullet(anchor, "WBS Viewer: Tree-style layout showing project, phases, activities, and tasks with expandable/collapsible nodes.")
    add_bullet(anchor, "Task Board: Kanban-style columns for todo, in_progress, and done tasks, with assignment and status update controls.")
    add_bullet(anchor, "Schedule View: Gantt-style chart showing task start dates, end dates, dependencies, and critical tasks.")
    add_bullet(anchor, "Critical Path View: Highlighted list or graph of zero-slack tasks that directly affect project duration.")
    add_placeholder(
        anchor,
        "Figure 9: Frontend Interface Design for ScaffoldPM",
        "Create a UI wireframe showing a ScaffoldPM dashboard with sidebar or top navigation, project generation form, document upload field, team member input area, WBS tree panel, task table/Kanban board, schedule view, and critical path summary.",
    )

    add_before(anchor, "Algorithm Details", style="Heading 2")
    add_before(anchor, "This section provides detailed algorithms for the core ScaffoldPM processing pipeline, reflecting the backend implementation.", justify=True)

    add_before(anchor, "Document Extraction Algorithm", style="Heading 3")
    add_before(anchor, "The document extraction algorithm converts uploaded scope documents into clean text. It supports PDF, DOC/DOCX, and TXT files. For PDF files, pdfplumber is used first, and PyMuPDF is used as a fallback when text extraction fails. DOC/DOCX files are processed using python-docx, while TXT files are decoded using UTF-8 or Latin-1 fallback.", justify=True)
    add_number(anchor, "Identify file extension and content type.")
    add_number(anchor, "Select the appropriate extractor for PDF, DOC/DOCX, or TXT.")
    add_number(anchor, "Extract page text, paragraph text, or raw text depending on the file type.")
    add_number(anchor, "Remove repeated headers, repeated footers, page numbers, and extra whitespace.")
    add_number(anchor, "Return cleaned scope text to the project generation workflow.")

    add_before(anchor, "WBS Generation Algorithm", style="Heading 3")
    add_before(anchor, "The WBS generation algorithm uses a large language model through LangChain to transform cleaned scope text into a structured project hierarchy. The output is parsed using Pydantic models so that phases, activities, tasks, dependencies, task types, and required skills are returned in a predictable schema.", justify=True)
    add_number(anchor, "Prepare project metadata, team member data, project dates, and extracted scope text.")
    add_number(anchor, "Build a prompt requesting 4 to 6 phases, 3 to 6 activities per phase, and 4 to 8 tasks per activity.")
    add_number(anchor, "Invoke the LLM and request valid JSON that matches the expected schema.")
    add_number(anchor, "Parse the LLM output into structured Pydantic objects.")
    add_number(anchor, "Generate internal UUIDs for project, phases, activities, and tasks.")
    add_number(anchor, "Resolve dependency names from task titles into internal task IDs.")

    add_before(anchor, "Semantic Task Matching Algorithm", style="Heading 3")
    add_before(anchor, "The task matching algorithm assigns tasks to team members based on semantic similarity, experience, task complexity, and workload distribution. SentenceTransformer embeddings are generated for task profiles and team member profiles.", justify=True)
    add_before(anchor, "Match score = (semantic_similarity * experience_multiplier) / load_penalty.", justify=True)
    add_number(anchor, "Create member profile text from role, skills, and experience.")
    add_number(anchor, "Create task profile text from title, description, task type, and required skills.")
    add_number(anchor, "Convert profiles into normalized embeddings using all-MiniLM-L6-v2.")
    add_number(anchor, "Compute similarity between each task and each team member.")
    add_number(anchor, "Adjust score using experience multiplier and workload penalty.")
    add_number(anchor, "Assign each task to the member with the highest adjusted score.")

    add_before(anchor, "Duration Estimation Algorithm", style="Heading 3")
    add_before(anchor, "The duration estimation module first attempts to load a trained XGBoost regression model. If the model file is available, it predicts estimated hours from task features. If the model is unavailable, the system uses fallback heuristic estimation.", justify=True)
    add_table_before(anchor, ["Feature", "Description"], [
        ["task_type", "Numeric value representing feature, bugfix, documentation, research, or other"],
        ["description_words", "Number of words in task description"],
        ["skills_count", "Number of required skills"],
        ["experience_value", "Numeric value for junior, mid, or senior assignee"],
        ["hours_estimate", "Heuristic base estimate used as a model feature"],
    ], [1.8, 4.5])
    add_before(anchor, "Heuristic estimate = base task hours adjusted by experience factor, required skills factor, and description length factor. The value is clamped between 2 and 40 hours and rounded to the nearest 0.5 hour.", justify=True)

    add_before(anchor, "Scheduling and Critical Path Algorithm", style="Heading 3")
    add_before(anchor, "The scheduling algorithm assigns start and end dates to tasks using dependencies, working days, and team member availability. Estimated hours are converted into duration days using 8 hours per working day. Weekends are skipped.", justify=True)
    add_number(anchor, "Create a task map and dependency map.")
    add_number(anchor, "Place dependency-free tasks in a ready queue.")
    add_number(anchor, "Select ready tasks by earliest feasible start date.")
    add_number(anchor, "Calculate task end date by adding working days.")
    add_number(anchor, "Update assigned member availability after each scheduled task.")
    add_number(anchor, "Repeat until all schedulable tasks are assigned dates.")
    add_before(anchor, "After scheduling, the CPM analyzer builds a directed graph where tasks are nodes and dependencies are edges. It calculates earliest start, earliest finish, latest start, latest finish, and slack. Tasks with zero slack are marked as critical.", justify=True)

    add_before(anchor, "Model Parameters and Configuration", style="Heading 3")
    add_table_before(anchor, ["Parameter", "Value", "Description"], [
        ["Backend framework", "FastAPI", "REST API framework"],
        ["Database", "SQLite with aiosqlite", "Local relational persistence"],
        ["LLM orchestration", "LangChain + OpenAI API", "WBS generation from scope text"],
        ["Embedding model", "all-MiniLM-L6-v2", "Semantic task-member matching"],
        ["Duration model", "XGBoost Regressor", "Task effort estimation"],
        ["Fallback duration", "Heuristic rules", "Used when trained model is unavailable"],
        ["Working hours/day", "8 hours", "Conversion from estimated hours to working days"],
        ["Calendar rule", "Monday-Friday", "Weekend days are skipped"],
        ["Graph library", "NetworkX", "Critical path and topological ordering"],
        ["Default database URL", "sqlite+aiosqlite:///./scaffoldpm.db", "Development database location"],
        ["Allowed frontend origin", "http://localhost:3000", "CORS origin for future Next.js frontend"],
    ], [1.8, 1.8, 2.7])

    add_before(
        anchor,
        "This completes the system design specification for ScaffoldPM. The modular FastAPI backend, document extraction workflow, LLM-based WBS generation, semantic task matching, duration prediction, dependency-aware scheduling, critical path analysis, and relational persistence design provide a clear blueprint for implementation and future frontend integration.",
        justify=True,
    )


def main() -> None:
    doc = Document(REPORT)
    update_figures(doc)
    anchor = remove_between(doc, "CHAPTER 4 - SYSTEM DESIGN", "REFERENCES", after_text="CHAPTER 3 - SYSTEM ANALYSIS")
    add_chapter_4(anchor)
    doc.save(OUT_REPORT)
    print(OUT_REPORT)


if __name__ == "__main__":
    main()
