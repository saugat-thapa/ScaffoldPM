from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


REPORT = "ScaffoldPM_Final_Year_Project_Report_Chapters_1_to_4.docx"


def set_run_font(run, size=11, bold=False, italic=False, color=None) -> None:
    run.font.name = "Calibri"
    if run._element.rPr is not None:
        run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def clear_runs(paragraph, text: str) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    run = paragraph.add_run(text)
    set_run_font(run)


def add_before(anchor, text: str, style: str | None = None, justify: bool = False, italic: bool = False):
    paragraph = anchor.insert_paragraph_before(text)
    if style:
        paragraph.style = style
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    for run in paragraph.runs:
        set_run_font(run, italic=italic)
    return paragraph


def add_figure_placeholder(anchor, title: str, description: str) -> None:
    title_p = add_before(anchor, title)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title_p.runs:
        set_run_font(run, size=10, bold=True, color="1F4D78")

    desc_p = add_before(anchor, f"[Insert {title.lower()} here. Suggested image description: {description}]", italic=True)
    desc_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for run in desc_p.runs:
        set_run_font(run, size=10, italic=True, color="555555")


def remove_elements_between(doc: Document, start_text: str, end_text: str, after_text: str | None = None):
    paragraphs = list(doc.paragraphs)
    after_idx = -1
    if after_text is not None:
        after_idx = next(i for i, p in enumerate(paragraphs) if p.text.strip() == after_text)

    start_idx_para = next(
        i for i, p in enumerate(paragraphs)
        if i > after_idx and p.text.strip() == start_text
    )
    end_idx_para = next(
        i for i, p in enumerate(paragraphs)
        if i > start_idx_para and p.text.strip() == end_text
    )
    start_para = paragraphs[start_idx_para]
    end_para = paragraphs[end_idx_para]

    body = doc._body._element
    children = list(body)
    start_idx = children.index(start_para._element)
    end_idx = children.index(end_para._element)

    for element in children[start_idx:end_idx]:
        body.remove(element)

    return end_para


def update_list_of_figures(doc: Document) -> None:
    paragraphs = list(doc.paragraphs)
    start = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "LIST OF FIGURES")
    end = next(i for i, p in enumerate(paragraphs) if i > start and p.text.strip() == "LIST OF TABLES")
    start_para = paragraphs[start]
    end_para = paragraphs[end]

    body = doc._body._element
    children = list(body)
    start_child = children.index(start_para._element)
    end_child = children.index(end_para._element)
    for element in children[start_child + 1:end_child]:
        body.remove(element)

    for item in [
        "Figure 1: Use Case Diagram for ScaffoldPM",
        "Figure 2: Gantt Chart for ScaffoldPM Development",
        "Figure 3: Class Diagram for ScaffoldPM",
        "Figure 4: Sequence Diagram for Project Plan Generation",
        "Figure 5: Activity Diagram for AI-Assisted Project Planning Pipeline",
        "Figure 6: State Diagram for Project Generation Request Lifecycle",
        "Figure 7: System Flowchart for ScaffoldPM",
        "Figure 8: Form and Report Design for ScaffoldPM",
        "Figure 9: Interface and Dialogue Design for ScaffoldPM",
    ]:
        p = end_para.insert_paragraph_before(item)
        for run in p.runs:
            set_run_font(run)


def update_chapter4_figure_numbers(doc: Document) -> None:
    replacements = {
        "Figure 5: System Flowchart for ScaffoldPM": "Figure 7: System Flowchart for ScaffoldPM",
        "Figure 6: Form and Report Design for ScaffoldPM": "Figure 8: Form and Report Design for ScaffoldPM",
        "Figure 7: Interface and Dialogue Design for ScaffoldPM": "Figure 9: Interface and Dialogue Design for ScaffoldPM",
    }
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text in replacements:
            clear_runs(paragraph, replacements[text])


def update_report_organization(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == "Chapter 3 describes system analysis, requirements, feasibility analysis, and structured analysis models.":
            clear_runs(
                paragraph,
                "Chapter 3 describes system analysis, requirements, feasibility analysis, and object-oriented analysis models.",
            )


def main() -> None:
    doc = Document(REPORT)
    update_list_of_figures(doc)
    update_chapter4_figure_numbers(doc)
    update_report_organization(doc)

    ch4_anchor = remove_elements_between(
        doc,
        "Analysis (Structured Approach)",
        "CHAPTER 4 - SYSTEM DESIGN",
        after_text="CHAPTER 3 - SYSTEM ANALYSIS",
    )

    add_before(ch4_anchor, "Object-Oriented Analysis", style="Heading 2")
    add_before(
        ch4_anchor,
        "Object-Oriented Analysis (OOA) is used to model ScaffoldPM through UML-based conceptual diagrams. It helps describe the static structure of the system, the interaction between system components, the internal workflow of project plan generation, and the lifecycle of a generated planning request. Since ScaffoldPM is designed as a modular backend system, OOA is useful for understanding how entities such as Project, Phase, Activity, Task, TeamMember, and WBSResponse interact with service modules such as DocumentExtractor, WBSGenerator, TaskMatcher, DurationPredictor, Scheduler, CPMAnalyzer, and ProjectRepository.",
        justify=True,
    )
    add_before(
        ch4_anchor,
        "The object-oriented view of ScaffoldPM focuses on four major UML diagrams: class diagram, sequence diagram, activity diagram, and state diagram. The class diagram shows the static structure of the system. The sequence diagram explains how a project generation request flows through different components. The activity diagram presents the complete AI-assisted project planning pipeline. The state diagram shows the lifecycle of a project generation request from upload to completion or failure.",
        justify=True,
    )

    add_before(ch4_anchor, "Class Diagram", style="Heading 3")
    add_before(
        ch4_anchor,
        "The class diagram identifies the core classes, data models, services, and relationships used in ScaffoldPM. The Project class represents the main project plan and contains project metadata such as project name, description, start date, end date, total working days, and created date. A Project contains one or more Phase objects, and each Phase contains one or more Activity objects. Each Activity contains multiple Task objects. The Task class stores task-specific information such as title, description, task type, status, estimated hours, start date, end date, dependencies, required skills, assigned member, slack, and critical path flag.",
        justify=True,
    )
    add_before(
        ch4_anchor,
        "The TeamMember class represents members available for the project. Each team member has an id, name, role, skills, and experience level. Tasks can be assigned to team members based on semantic skill matching. The WBSResponse class represents the generated project plan returned by the system, including the project object, optional summary, and total duration in working days.",
        justify=True,
    )
    add_before(
        ch4_anchor,
        "The service classes represent the main processing modules of ScaffoldPM. DocumentExtractor extracts and cleans text from uploaded documents. WBSGenerator uses the extracted scope text and project metadata to generate phases, activities, and tasks. TaskMatcher assigns tasks to suitable team members. DurationPredictor estimates the effort required for each task. Scheduler assigns start and end dates based on dependencies and member availability. CPMAnalyzer calculates slack and identifies critical tasks. ProjectRepository handles database operations such as saving, fetching, listing, and updating project data.",
        justify=True,
    )
    add_figure_placeholder(
        ch4_anchor,
        "Figure 3: Class Diagram for ScaffoldPM",
        "Create a UML class diagram showing Project, Phase, Activity, Task, TeamMember, WBSResponse, DocumentExtractor, WBSGenerator, TaskMatcher, DurationPredictor, Scheduler, CPMAnalyzer, and ProjectRepository. Show Project-to-Phase, Phase-to-Activity, Activity-to-Task, TeamMember-to-Task, Project-to-TeamMember, Task self-dependency, and service dependency relationships.",
    )

    add_before(ch4_anchor, "Sequence Diagram: Project Plan Generation", style="Heading 3")
    add_before(
        ch4_anchor,
        "The sequence diagram describes the interaction between the user and ScaffoldPM components during project plan generation. The process begins when the Project Manager uploads a scope document and submits project metadata such as project name, dates, and team member details. The frontend or API client sends a multipart request to the FastAPI backend.",
        justify=True,
    )
    add_before(
        ch4_anchor,
        "The ProjectRouter validates the payload and forwards the uploaded file to the DocumentExtractor. The DocumentExtractor extracts and cleans the text from the scope document. The cleaned scope text is passed to the WBSGenerator, which uses the LLM service to generate a structured work breakdown structure. After the WBS is generated, the TaskMatcher assigns tasks to suitable team members using semantic similarity. The DurationPredictor estimates task effort, the Scheduler calculates task start and end dates, and the CPMAnalyzer identifies critical tasks. Finally, the ProjectRepository stores the generated project plan in the database and the API returns the complete project plan to the user.",
        justify=True,
    )
    add_before(
        ch4_anchor,
        "Participants: Project Manager, Frontend/API Client, FastAPI App, ProjectRouter, DocumentExtractor, WBSGenerator, OpenAI LLM Service, TaskMatcher, DurationPredictor, Scheduler, CPMAnalyzer, ProjectRepository, SQLite Database.",
        justify=True,
    )
    add_figure_placeholder(
        ch4_anchor,
        "Figure 4: Sequence Diagram for Project Plan Generation",
        "Create a UML sequence diagram with lifelines for Project Manager, Frontend/API Client, FastAPI App, ProjectRouter, DocumentExtractor, WBSGenerator, OpenAI LLM Service, TaskMatcher, DurationPredictor, Scheduler, CPMAnalyzer, ProjectRepository, and SQLite Database. Show the POST /api/project/generate-wbs flow from upload to saved project response, including validation and common error responses.",
    )

    add_before(ch4_anchor, "Activity Diagram: AI-Assisted Project Planning Pipeline", style="Heading 3")
    add_before(
        ch4_anchor,
        "The activity diagram represents the complete workflow followed by ScaffoldPM during project plan generation. The process starts when the user uploads a scope document and submits the project payload. The system validates the request and checks whether the uploaded document type is supported. If the request is invalid or the file type is unsupported, the system returns an error response.",
        justify=True,
    )
    add_before(
        ch4_anchor,
        "If the input is valid, the document text is extracted and cleaned. The cleaned text is used by the WBS generation module to produce phases, activities, tasks, required skills, and dependencies. The system then resolves dependencies, assigns tasks to team members, estimates duration, schedules tasks, analyzes the critical path, stores the project, and returns the generated plan. This activity flow shows the internal decision points and processing stages of ScaffoldPM.",
        justify=True,
    )
    add_figure_placeholder(
        ch4_anchor,
        "Figure 5: Activity Diagram for AI-Assisted Project Planning Pipeline",
        "Create a UML activity diagram showing request validation, file type checking, document extraction, text cleaning, LLM-based WBS generation, dependency resolution, semantic task assignment, duration estimation with ML-or-heuristic decision, scheduling, critical path analysis, database saving, success response, and error paths.",
    )

    add_before(ch4_anchor, "State Diagram: Project Generation Request Lifecycle", style="Heading 3")
    add_before(
        ch4_anchor,
        "The state diagram shows the lifecycle of a project generation request in ScaffoldPM. A request begins in the idle state before the user uploads a document and submits project details. After submission, the request moves to validation. If validation fails, the request enters an error state. If validation succeeds, the request moves through document extraction, WBS generation, task assignment, duration estimation, scheduling, critical path analysis, and database persistence.",
        justify=True,
    )
    add_before(
        ch4_anchor,
        "When all processing steps complete successfully, the request enters the completed state and the generated project plan is returned to the user. If an error occurs during extraction, WBS generation, database saving, or any major processing step, the request moves to a failed state with an appropriate HTTP error response. This state model helps describe how ScaffoldPM handles both successful and unsuccessful project generation workflows.",
        justify=True,
    )
    add_figure_placeholder(
        ch4_anchor,
        "Figure 6: State Diagram for Project Generation Request Lifecycle",
        "Create a UML state diagram with states Idle, Submitted, Validating Request, Invalid Request, Extracting Document Text, Generating WBS, Resolving Dependencies, Assigning Tasks, Estimating Durations, Scheduling Tasks, Analyzing Critical Path, Saving Project, Completed, and Failed. Include transitions for success and HTTP error cases.",
    )

    doc.save(REPORT)


if __name__ == "__main__":
    main()
