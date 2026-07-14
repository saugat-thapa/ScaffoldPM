from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = Path("ScaffoldPM_Final_Year_Project_Report_Chapters_1_to_4.docx")


TITLE = "ScaffoldPM"
SUBTITLE = "AI-Driven Project Planning and Optimization System"
COLLEGE = "ST. XAVIER'S COLLEGE"
PROGRAM = "Bachelor of Science in Computer Science and Information Technology (B.Sc.CSIT)"
UNIVERSITY = "TRIBHUVAN UNIVERSITY\nINSTITUTE OF SCIENCE AND TECHNOLOGY"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths_inches: list[float]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_inches):
            cell = row.cells[idx]
            cell.width = Inches(width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def set_font(run, name="Calibri", size=11, bold=False, italic=False, color=None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def setup_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_style in ["List Bullet", "List Number"]:
        style = styles[list_style]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.15


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)


def add_centered(doc: Document, text: str, size=12, bold=False, space_after=6) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    for idx, line in enumerate(text.split("\n")):
        if idx:
            p.add_run().add_break()
        r = p.add_run(line)
        set_font(r, size=size, bold=bold)


def add_para(doc: Document, text: str, style=None, justify=False) -> None:
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    set_font(r)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        add_para(doc, item, style="List Bullet")


def add_numbers(doc: Document, items: list[str]) -> None:
    for item in items:
        add_para(doc, item, style="List Number")


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    set_font(r, size=10, italic=True, color="555555")


def add_figure_placeholder(doc: Document, title: str, description: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_width(table, [6.3])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F6F9")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_font(r, size=11, bold=True, color="1F4D78")
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run("Image placeholder description: " + description)
    set_font(r2, size=10, italic=True, color="333333")


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table, widths)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, header in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, "E8EEF5")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        set_font(r, size=10, bold=True, color="0B2545")
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_font(r, size=10)
    doc.add_paragraph()


def add_front_matter(doc: Document) -> None:
    add_centered(doc, UNIVERSITY, size=14, bold=True, space_after=18)
    add_centered(doc, "Final Year Project Report on", size=12, space_after=8)
    add_centered(doc, TITLE, size=18, bold=True, space_after=2)
    add_centered(doc, SUBTITLE, size=12, bold=True, space_after=18)
    add_centered(doc, "For the partial fulfillment of the requirements of the degree of", size=11)
    add_centered(doc, PROGRAM, size=12, bold=True, space_after=8)
    add_centered(doc, "awarded by Tribhuvan University", size=11, space_after=18)
    add_centered(doc, "Submitted By:", size=12, bold=True, space_after=4)
    add_centered(doc, "[Student Name 1] ([Roll Number])\n[Student Name 2] ([Roll Number])", size=11, space_after=14)
    add_centered(doc, COLLEGE, size=12, bold=True, space_after=2)
    add_centered(doc, "Maitighar, Kathmandu, Nepal", size=11, space_after=14)
    add_centered(doc, "Submitted To:", size=12, bold=True, space_after=4)
    add_centered(doc, "Office of the Dean\nInstitute of Science and Technology\nTribhuvan University, Kathmandu", size=11, space_after=18)
    add_centered(doc, "June 2026", size=11)
    doc.add_page_break()

    doc.add_heading("SUPERVISOR RECOMMENDATION", level=1)
    add_para(doc, "This is to certify that the final year project entitled \"ScaffoldPM: AI-Driven Project Planning and Optimization System\" is an academic work completed by [Student Name 1] and [Student Name 2] under my supervision. The project has been prepared for the partial fulfillment of the requirements for the degree of Bachelor of Science in Computer Science and Information Technology.", justify=True)
    add_para(doc, "The information presented in this project report has not been submitted earlier to the best of my knowledge. I recommend this report for evaluation.", justify=True)
    doc.add_paragraph("\n\n\nSignature of the Supervisor\nName: [Supervisor Name]\nDesignation: [Designation]\nDate: June 2026")
    doc.add_page_break()

    doc.add_heading("CERTIFICATE OF APPROVAL", level=1)
    add_para(doc, "The undersigned certify that they have read and recommended to the Department of Computer Science, St. Xavier's College for acceptance, a final year project report entitled \"ScaffoldPM: AI-Driven Project Planning and Optimization System\" submitted by [Student Name 1] and [Student Name 2] in partial fulfillment of the requirements for the degree of Bachelor of Science in Computer Science and Information Technology.", justify=True)
    doc.add_paragraph("\n\n.................................\nProject Supervisor\n\n.................................\nInternal Examiner\n\n.................................\nHead of Department\n\n.................................\nExternal Examiner")
    doc.add_page_break()

    doc.add_heading("ACKNOWLEDGEMENT", level=1)
    for text in [
        "We would like to express our sincere gratitude to the Department of Computer Science, St. Xavier's College, Maitighar, Kathmandu, for providing us with the opportunity to undertake this final year project.",
        "We are deeply thankful to our respected supervisor for continuous guidance, suggestions, and encouragement throughout the development of ScaffoldPM. The feedback received during planning, implementation, and documentation helped us improve both the technical quality and academic presentation of the project.",
        "We also extend our thanks to the faculty members, friends, and all individuals who directly or indirectly supported us during the completion of this work.",
    ]:
        add_para(doc, text, justify=True)
    doc.add_paragraph("\n[Student Name 1]\n[Student Name 2]")
    doc.add_page_break()

    doc.add_heading("ABSTRACT", level=1)
    for text in [
        "ScaffoldPM is an AI-driven project planning and optimization system designed to automate the generation of structured project plans from unstructured scope documents. The system accepts a project scope document and basic project metadata, extracts readable text, and uses a large language model to generate a work breakdown structure containing phases, activities, tasks, required skills, and dependencies.",
        "After generating the WBS, ScaffoldPM assigns tasks to suitable team members using semantic skill matching, estimates task duration using an XGBoost-based machine learning model or fallback heuristic logic, schedules tasks according to dependencies and team availability, and identifies critical tasks through critical path analysis. The backend is implemented using FastAPI, Pydantic, SQLAlchemy, SQLite, LangChain, OpenAI API, Sentence Transformers, XGBoost, and NetworkX.",
        "The project demonstrates how artificial intelligence, natural language processing, machine learning, and project management principles can be combined to reduce manual planning effort and improve project organization.",
    ]:
        add_para(doc, text, justify=True)
    add_para(doc, "Keywords: Artificial Intelligence, Project Management, Work Breakdown Structure, Critical Path Method, FastAPI, XGBoost, Sentence Transformers.")
    doc.add_page_break()

    doc.add_heading("TABLE OF CONTENTS", level=1)
    toc = [
        "SUPERVISOR RECOMMENDATION",
        "CERTIFICATE OF APPROVAL",
        "ACKNOWLEDGEMENT",
        "ABSTRACT",
        "LIST OF ABBREVIATIONS",
        "LIST OF FIGURES",
        "LIST OF TABLES",
        "CHAPTER 1 - INTRODUCTION",
        "CHAPTER 2 - BACKGROUND STUDY AND LITERATURE REVIEW",
        "CHAPTER 3 - SYSTEM ANALYSIS",
        "CHAPTER 4 - SYSTEM DESIGN",
        "REFERENCES",
    ]
    for item in toc:
        add_para(doc, item)
    doc.add_page_break()

    doc.add_heading("LIST OF ABBREVIATIONS", level=1)
    add_table(doc, ["Abbreviation", "Full Form"], [
        ["AI", "Artificial Intelligence"],
        ["API", "Application Programming Interface"],
        ["CPM", "Critical Path Method"],
        ["CRUD", "Create, Read, Update, Delete"],
        ["DFD", "Data Flow Diagram"],
        ["ERD", "Entity Relationship Diagram"],
        ["LLM", "Large Language Model"],
        ["ML", "Machine Learning"],
        ["NLP", "Natural Language Processing"],
        ["ORM", "Object Relational Mapping"],
        ["WBS", "Work Breakdown Structure"],
    ], [1.6, 4.7])
    doc.add_page_break()

    doc.add_heading("LIST OF FIGURES", level=1)
    for fig in [
        "Figure 1: Use Case Diagram for ScaffoldPM",
        "Figure 2: Gantt Chart / Project Schedule for ScaffoldPM Development",
        "Figure 3: Entity Relationship Diagram for ScaffoldPM",
        "Figure 4: Data Flow Diagram for ScaffoldPM",
        "Figure 5: System Flowchart for ScaffoldPM",
        "Figure 6: Form and Report Design for ScaffoldPM",
        "Figure 7: Interface and Dialogue Design for ScaffoldPM",
    ]:
        add_para(doc, fig)
    doc.add_heading("LIST OF TABLES", level=1)
    for tbl in [
        "Table 1: Functional Requirements of ScaffoldPM",
        "Table 2: Non-Functional Requirements of ScaffoldPM",
        "Table 3: Hardware Requirements",
        "Table 4: Software Requirements",
        "Table 5: Entity to Relation Mapping",
    ]:
        add_para(doc, tbl)
    doc.add_page_break()


def add_chapter_1(doc: Document) -> None:
    doc.add_heading("CHAPTER 1 - INTRODUCTION", level=1)
    doc.add_heading("Introduction", level=2)
    for text in [
        "Project planning is one of the most important activities in software engineering and information system development. Before a project can be implemented successfully, the project manager must understand the scope, divide the work into smaller components, estimate the time required for each task, assign suitable team members, and identify activities that can delay the whole project. According to the PMBOK Guide, project management involves coordinated planning and control of scope, schedule, resources, and related activities [1].",
        "In many organizations, the initial planning process is still performed manually. The project manager reads a requirement document, prepares phases and tasks, estimates duration, and assigns work based on personal judgment. This process consumes time and can introduce human error, especially when the requirement document is long or unclear. Inaccurate software effort estimation is a known challenge because wrong estimates can affect cost, schedule, and delivery reliability [8].",
        "ScaffoldPM is proposed as an AI-driven project planning and optimization system that assists project managers in converting an uploaded scope document into a structured project plan. It extracts text from PDF, DOC/DOCX, or TXT documents, generates a work breakdown structure using a large language model, assigns tasks using semantic skill matching, estimates duration using a machine learning model or heuristic fallback, schedules tasks, identifies critical path tasks, and stores the final plan in a database.",
    ]:
        add_para(doc, text, justify=True)

    doc.add_heading("Problem Statement", level=2)
    for text in [
        "Traditional project planning requires a project manager to manually read the scope document and identify phases, activities, tasks, dependencies, required skills, and schedules. This process is repetitive and depends heavily on the experience of the project manager. Important tasks may be missed, dependencies may be assigned incorrectly, and duration estimates may become inconsistent.",
        "Another major issue is task assignment. When team members have different roles, skills, and experience levels, assigning work manually can lead to imbalance. Some members may receive too many tasks while other suitable members remain underused. Similarly, without proper dependency analysis, project managers may fail to identify critical tasks that directly affect the project deadline.",
        "Therefore, there is a need for a system that can automatically generate a structured project plan from a scope document and support project managers with task assignment, effort estimation, scheduling, and critical path analysis.",
    ]:
        add_para(doc, text, justify=True)

    doc.add_heading("Objectives", level=2)
    add_para(doc, "The main objective of this project is to design and implement an AI-driven backend system that generates and optimizes project plans from uploaded scope documents.")
    add_numbers(doc, [
        "To extract readable text from uploaded PDF, DOC/DOCX, and TXT project scope documents.",
        "To generate a detailed work breakdown structure containing phases, activities, and tasks.",
        "To identify task types, required skills, and task dependencies.",
        "To assign tasks to team members based on semantic similarity, skill match, experience, and workload.",
        "To estimate task duration using an XGBoost model or heuristic fallback logic.",
        "To schedule tasks according to dependencies, working days, and team member availability.",
        "To calculate slack and identify critical path tasks.",
        "To store generated project plans in a relational database and expose them through REST APIs.",
    ])

    doc.add_heading("Scope and Limitations", level=2)
    doc.add_heading("Scope", level=3)
    add_para(doc, "The project focuses on the backend implementation of an AI-assisted project planning system. The major scope of ScaffoldPM includes the following:")
    add_bullets(doc, [
        "Uploading and processing project scope documents.",
        "Extracting and cleaning text from PDF, DOC/DOCX, and TXT files.",
        "Generating WBS data through an LLM-powered module.",
        "Managing projects, phases, activities, tasks, and team members.",
        "Assigning tasks using semantic embeddings and workload balancing.",
        "Estimating task duration through a trained model or heuristic logic.",
        "Scheduling tasks while skipping weekends.",
        "Analyzing task dependencies using critical path method.",
        "Persisting project data in SQLite through SQLAlchemy ORM.",
        "Providing API endpoints for project generation, retrieval, task status update, and reassignment.",
    ])

    doc.add_heading("Limitations", level=3)
    add_bullets(doc, [
        "The quality of the generated WBS depends on the quality of the uploaded scope document and the LLM response.",
        "The current implementation requires an OpenAI API key for WBS generation.",
        "The frontend directory exists but a complete production frontend is not implemented in the current repository.",
        "The scheduler uses a greedy approach and may not always produce the globally optimal schedule.",
        "The system excludes weekends but does not currently consider public holidays, leave days, or part-time availability.",
        "Authentication, authorization, and role-based access control are not implemented.",
        "Skills and dependencies are stored as comma-separated text rather than fully normalized relational tables.",
    ])

    doc.add_heading("Development Methodology", level=2)
    for text in [
        "An iterative and incremental development methodology was followed for ScaffoldPM. The project was divided into independent modules such as document extraction, WBS generation, task matching, duration prediction, scheduling, critical path analysis, and database persistence. Each module was implemented and tested separately before being integrated into the main API workflow.",
        "This approach was suitable because the project combines multiple technologies and algorithms. Building the system incrementally allowed the document extraction module to be validated first, followed by WBS generation, task assignment, duration estimation, scheduling, and database operations. The modular design also supports future extension such as frontend integration, authentication, deployment, and improved optimization.",
    ]:
        add_para(doc, text, justify=True)

    doc.add_heading("Report Organization", level=2)
    add_para(doc, "The report is organized into the following chapters:")
    add_bullets(doc, [
        "Chapter 1 introduces the project, problem statement, objectives, scope, limitations, development methodology, and report organization.",
        "Chapter 2 presents the background study and literature review related to project planning, WBS, CPM, effort estimation, embeddings, and technologies used.",
        "Chapter 3 describes system analysis, requirements, feasibility analysis, and structured analysis models.",
        "Chapter 4 explains the system design, including architecture, database design, interface design, and algorithm descriptions.",
    ])


def add_chapter_2(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("CHAPTER 2 - BACKGROUND STUDY AND LITERATURE REVIEW", level=1)
    doc.add_heading("Background Study", level=2)
    for text in [
        "Project planning involves defining the work required to complete a project and organizing it into manageable components. A Work Breakdown Structure is widely used to decompose project scope into smaller deliverables and work packages [2]. In software projects, this decomposition is useful because it allows project teams to understand the required work, estimate effort, assign responsibilities, and track progress.",
        "Critical Path Method is another important project management technique. It represents project activities and their dependencies and identifies the sequence of tasks that controls the total project duration [4]. Tasks on the critical path have zero slack, meaning any delay in those tasks can delay the entire project. ScaffoldPM uses this idea to mark tasks that require special attention from the project manager.",
        "Software effort estimation is a difficult problem because development effort depends on task complexity, developer experience, technology stack, requirement clarity, and many other factors. Traditional approaches such as COCOMO provided early structured models for effort estimation [7]. More recent studies have explored machine learning techniques for improving software effort estimation accuracy when historical project data is available [9].",
        "Natural language processing and large language models provide new opportunities for transforming unstructured requirement documents into structured outputs. ScaffoldPM uses an LLM to generate phases, activities, tasks, required skills, and dependencies from scope text. It also uses semantic sentence embeddings for matching tasks with team members. Sentence-BERT shows that sentence embeddings can represent sentence meaning in a way that supports semantic similarity comparison [5].",
    ]:
        add_para(doc, text, justify=True)

    doc.add_heading("Literature Review", level=2)
    literature_rows = [
        ["Project Management Institute (2021)", "Provides general project management principles including planning, scope, schedule, and resource coordination.", "Supports the overall project management foundation of ScaffoldPM."],
        ["Project Management Institute (2019)", "Explains WBS as a structured decomposition of project scope into manageable work components.", "Supports the WBS-based phase, activity, and task structure."],
        ["Hans (2013)", "Discusses WBS as a tool for software project scope verification.", "Justifies using WBS in software project planning."],
        ["Kelley and Walker (1959)", "Introduces critical-path planning and scheduling.", "Supports dependency analysis and critical task identification."],
        ["Boehm (1981)", "Presents foundational work in software engineering economics and effort estimation.", "Supports the need for systematic software effort estimation."],
        ["Reimers and Gurevych (2019)", "Introduces Sentence-BERT for semantically meaningful sentence embeddings.", "Supports semantic task-team matching."],
        ["Chen and Guestrin (2016)", "Introduces XGBoost as a scalable tree boosting system.", "Supports the task duration prediction model."],
    ]
    add_table(doc, ["Source", "Main Contribution", "Use in ScaffoldPM"], literature_rows, [1.65, 2.45, 2.2])

    doc.add_heading("Existing System", level=2)
    for text in [
        "Many project management tools provide task boards, Gantt charts, calendars, and collaboration features. However, these systems often require the project manager to manually create tasks, estimate durations, define dependencies, and assign team members. The planning output depends on the user's ability to interpret requirements and convert them into a structured plan.",
        "Existing manual systems also have limited intelligence in matching tasks to team members. A manager may know the team personally, but in larger teams or unfamiliar projects, it becomes difficult to maintain fair workload distribution and skill-based assignment. Similarly, traditional systems may show schedules but do not automatically generate the plan from the original scope document.",
    ]:
        add_para(doc, text, justify=True)

    doc.add_heading("Proposed System", level=2)
    add_para(doc, "ScaffoldPM proposes an AI-assisted backend system that reduces manual effort by converting the uploaded scope document into a structured, scheduled, and stored project plan. The proposed system integrates the following components:")
    add_bullets(doc, [
        "Document extraction for reading PDF, DOC/DOCX, and TXT documents.",
        "LLM-based WBS generation for producing phases, activities, tasks, dependencies, and required skills.",
        "Semantic task assignment using SentenceTransformer embeddings.",
        "Duration estimation using XGBoost model or heuristic fallback.",
        "Dependency-aware scheduling using working-day calculations.",
        "Critical path analysis using directed graph processing.",
        "Database persistence and REST APIs for project management operations.",
    ])

    doc.add_heading("Technology Review", level=2)
    tech_rows = [
        ["FastAPI", "Python web framework used to build REST APIs with automatic documentation [10].", "API layer"],
        ["Pydantic", "Data validation and schema definition library [11].", "Request/response validation"],
        ["SQLAlchemy", "ORM used for database models and asynchronous operations [12].", "Database layer"],
        ["SQLite", "Lightweight serverless database [13].", "Local persistence"],
        ["LangChain", "Framework for connecting application logic with LLMs [14].", "WBS generation workflow"],
        ["OpenAI API", "LLM text generation and structured output support [15], [16].", "AI-generated WBS"],
        ["Sentence Transformers", "Library for sentence embeddings and semantic similarity [17].", "Task-member matching"],
        ["XGBoost", "Gradient boosting library for supervised learning [18].", "Duration prediction"],
        ["NetworkX", "Graph processing library with DAG utilities [19].", "Critical path analysis"],
    ]
    add_table(doc, ["Technology", "Description", "Role in System"], tech_rows, [1.4, 3.1, 1.8])


def add_chapter_3(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("CHAPTER 3 - SYSTEM ANALYSIS", level=1)
    doc.add_heading("System Analysis", level=2)
    for text in [
        "System analysis identifies the requirements, feasibility, users, processes, and data flow of the proposed system. ScaffoldPM is analyzed as a backend-centered project planning system that receives an uploaded document and project metadata, processes them through AI and scheduling modules, stores the result, and provides APIs for later retrieval and management.",
        "The primary actor of the system is the project manager or team lead. The user uploads a scope document, provides project dates and team member details, and receives a generated project plan. The generated plan contains tasks, assignments, estimated hours, schedules, and critical path information.",
    ]:
        add_para(doc, text, justify=True)

    doc.add_heading("Requirement Analysis", level=2)
    doc.add_heading("Functional Requirements", level=3)
    add_table(doc, ["ID", "Requirement"], [
        ["FR1", "The system shall allow the user to upload a project scope document."],
        ["FR2", "The system shall extract text from PDF, DOC/DOCX, and TXT documents."],
        ["FR3", "The system shall accept project name, start date, end date, and team member details."],
        ["FR4", "The system shall generate a WBS containing phases, activities, and tasks."],
        ["FR5", "The system shall identify task types, required skills, and dependencies."],
        ["FR6", "The system shall assign tasks to suitable team members."],
        ["FR7", "The system shall estimate task duration in hours."],
        ["FR8", "The system shall schedule tasks based on dependencies and availability."],
        ["FR9", "The system shall identify critical path tasks and slack values."],
        ["FR10", "The system shall save generated project plans in the database."],
        ["FR11", "The system shall list saved projects and fetch a project by ID."],
        ["FR12", "The system shall update task status and task assignment."],
    ], [0.8, 5.5])

    doc.add_heading("Non-Functional Requirements", level=3)
    add_table(doc, ["ID", "Requirement"], [
        ["NFR1", "The system should validate incoming data using defined schemas."],
        ["NFR2", "The system should provide structured JSON API responses."],
        ["NFR3", "The system should handle invalid documents and malformed payloads gracefully."],
        ["NFR4", "The backend should use modular architecture for maintainability."],
        ["NFR5", "The database operations should be asynchronous where possible."],
        ["NFR6", "The system should use fallback duration estimation if the ML model is unavailable."],
        ["NFR7", "The system should be extensible for future frontend, authentication, and deployment."],
    ], [0.8, 5.5])

    doc.add_heading("Use Case Analysis", level=2)
    add_para(doc, "The main use cases of ScaffoldPM are performed by the project manager. The project manager can generate a WBS, view stored projects, update task status, and reassign tasks.")
    add_figure_placeholder(doc, "Figure 1: Use Case Diagram for ScaffoldPM", "Create a UML use case diagram with one actor named Project Manager. Connect the actor to these use cases: Upload Scope Document, Enter Project Details, Generate WBS, View Project Plan, View Critical Path, List Saved Projects, Update Task Status, and Reassign Task. The system boundary should be labelled ScaffoldPM.")
    add_caption(doc, "Figure 1: Use Case Diagram for ScaffoldPM")

    doc.add_heading("Feasibility Analysis", level=2)
    doc.add_heading("Technical Feasibility Analysis", level=3)
    add_para(doc, "The project is technically feasible because the required technologies are available as open-source or API-based tools. FastAPI, SQLAlchemy, SQLite, Sentence Transformers, XGBoost, NetworkX, pdfplumber, PyMuPDF, and python-docx support the main backend operations. The WBS generation module depends on OpenAI API access through LangChain.")

    doc.add_heading("Hardware Requirements", level=3)
    add_table(doc, ["Component", "Minimum Requirement"], [
        ["Processor", "Intel Core i5 or equivalent"],
        ["RAM", "8 GB minimum, 16 GB recommended"],
        ["Storage", "2 GB free storage for source code, dependencies, dataset, and database"],
        ["Network", "Internet connection required for OpenAI API and model dependency download"],
    ], [1.8, 4.5])

    doc.add_heading("Software Requirements", level=3)
    add_table(doc, ["Software", "Purpose"], [
        ["Python", "Backend programming language"],
        ["FastAPI and Uvicorn", "API development and local server execution"],
        ["SQLite", "Local relational database"],
        ["OpenAI API Key", "LLM-based WBS generation"],
        ["Python libraries from requirements.txt", "Document extraction, validation, ML, NLP, scheduling, and graph analysis"],
        ["API testing tool", "Testing endpoints through Swagger UI or Postman"],
    ], [2.0, 4.3])

    doc.add_heading("Operational Feasibility", level=3)
    add_para(doc, "The system is operationally feasible because project managers already work with scope documents, project dates, teams, and tasks. ScaffoldPM fits into this workflow by reducing the manual effort required to prepare the initial plan. The generated plan can later be reviewed and modified by the manager.")

    doc.add_heading("Economic Feasibility", level=3)
    add_para(doc, "The system uses mainly open-source technologies, which reduces development cost. The main recurring cost is the LLM API usage required for WBS generation. For academic and prototype use, the system can run locally with SQLite and does not require paid database or hosting services.")

    doc.add_heading("Schedule Feasibility", level=3)
    add_para(doc, "The modular nature of the system makes it suitable for incremental development within a final year project schedule. Core modules can be implemented and tested independently, and future features such as frontend visualization and authentication can be added later.")
    add_figure_placeholder(doc, "Figure 2: Gantt Chart for ScaffoldPM Development", "Create a Gantt chart showing project phases across weeks: Requirement Study, Literature Review, Backend Setup, Document Extraction, WBS Generation, Task Matching, Duration Prediction, Scheduling and CPM, Database Integration, API Testing, Documentation, and Final Review.")
    add_caption(doc, "Figure 2: Gantt Chart for ScaffoldPM Development")

    doc.add_heading("Analysis (Structured Approach)", level=2)
    doc.add_heading("Data Modelling using ER Diagram", level=3)
    add_para(doc, "The main entities in ScaffoldPM are Project, Phase, Activity, Task, TeamMember, and ProjectMember. A project contains many phases, a phase contains many activities, and an activity contains many tasks. Projects and team members have a many-to-many relationship through the ProjectMember junction table.")
    add_figure_placeholder(doc, "Figure 3: Entity Relationship Diagram for ScaffoldPM", "Create an ERD with entities Project, Phase, Activity, Task, TeamMember, and ProjectMember. Show Project 1-to-many Phase, Phase 1-to-many Activity, Activity 1-to-many Task, TeamMember 1-to-many Task, and Project many-to-many TeamMember through ProjectMember. Include important attributes such as id, name, dates, status, estimated_hours, dependencies, skills_required, and slack.")
    add_caption(doc, "Figure 3: Data Modelling using ER Diagram for ScaffoldPM")

    doc.add_heading("Process Modeling using DFD Diagram", level=3)
    add_para(doc, "The process model of ScaffoldPM begins when the project manager submits a scope document and project metadata. The system extracts document text, generates WBS data, assigns team members, estimates duration, schedules tasks, performs critical path analysis, stores the project, and returns the generated plan.")
    add_figure_placeholder(doc, "Figure 4: Data Flow Diagram for ScaffoldPM", "Create a Level 1 DFD. External entity: Project Manager. Processes: Validate Request, Extract Document Text, Generate WBS, Assign Tasks, Estimate Duration, Schedule Tasks, Analyze Critical Path, Save Project, Retrieve/Update Project. Data store: SQLite Database. External service: OpenAI API. Show document, metadata, WBS, schedule, and stored project flows.")
    add_caption(doc, "Figure 4: DFD Diagram for ScaffoldPM")


def add_chapter_4(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("CHAPTER 4 - SYSTEM DESIGN", level=1)
    doc.add_heading("Design (Structured)", level=2)
    for text in [
        "The system design of ScaffoldPM follows a modular backend architecture. The API layer receives requests, validates data, and coordinates the project generation workflow. The module layer performs document extraction, WBS generation, matching, prediction, scheduling, and graph analysis. The data layer contains Pydantic schemas and SQLAlchemy database models. The persistence layer stores projects, phases, activities, tasks, and team members in SQLite.",
        "This modular design improves maintainability because each major responsibility is implemented in a separate module. Future components such as a frontend client, authentication service, notification service, Gantt chart visualization, and deployment configuration can be added without changing the complete backend structure.",
    ]:
        add_para(doc, text, justify=True)

    add_figure_placeholder(doc, "Figure 5: System Flowchart for ScaffoldPM", "Create a flowchart starting with Upload Scope Document and Project Payload. Then show Validate Payload, Extract Text, Generate WBS using LLM, Resolve Dependencies, Assign Tasks by Skill Matching, Estimate Durations, Schedule Tasks, Analyze Critical Path, Save to Database, and Return Project Plan. Add error paths for invalid file, invalid payload, LLM failure, and database failure.")
    add_caption(doc, "Figure 5: System Flowchart for ScaffoldPM")

    doc.add_heading("System Architecture", level=2)
    add_para(doc, "ScaffoldPM can be divided into the following architectural layers:")
    add_bullets(doc, [
        "Client Layer: A future frontend or API testing tool sends documents and project metadata to the backend.",
        "API Layer: FastAPI routes handle health checks, project generation, project listing, project retrieval, task status update, and task reassignment.",
        "AI and Processing Layer: The system extracts document text, generates WBS data, matches tasks, predicts duration, schedules tasks, and analyzes critical path.",
        "Data Access Layer: SQLAlchemy CRUD functions save, fetch, list, and update project data.",
        "Database Layer: SQLite stores all project-related entities.",
        "External AI Service: OpenAI API is used through LangChain for LLM-based WBS generation.",
    ])

    doc.add_heading("Frontend Layer (Client-Side)", level=3)
    add_para(doc, "The current repository contains an empty frontend directory. A future frontend can be developed using a web framework such as React or Next.js. The frontend would allow users to upload documents, enter team member details, view generated WBS, update tasks, and visualize schedules.")

    doc.add_heading("Backend Layer (Server-Side)", level=3)
    add_para(doc, "The backend is implemented using FastAPI. It exposes REST API endpoints under /api and /api/project. The main endpoint, /api/project/generate-wbs, accepts the uploaded file and project payload, then executes the complete planning workflow.")

    doc.add_heading("AI and Processing Modules", level=3)
    add_bullets(doc, [
        "Document extractor: Extracts and cleans text from uploaded files.",
        "WBS generator: Uses an LLM to create phases, activities, tasks, skills, and dependencies.",
        "Matching module: Assigns tasks to team members using semantic similarity and workload penalty.",
        "Prediction module: Estimates task duration using XGBoost or heuristic fallback.",
        "Scheduling module: Assigns start and end dates based on dependencies and member availability.",
        "CPM module: Calculates slack and marks critical tasks.",
    ])

    doc.add_heading("Data Persistence Layer", level=3)
    add_para(doc, "The persistence layer uses SQLAlchemy asynchronous sessions and SQLite. The database is initialized automatically when the FastAPI application starts. CRUD functions are responsible for saving complete project plans and updating task status or assignment.")

    doc.add_heading("Database Design: Transformation of ER to Relations and Normalization", level=2)
    add_table(doc, ["Entity/Table", "Important Attributes", "Description"], [
        ["projects", "id, name, description, start_date, end_date, total_working_days, summary, total_duration_days, created_at", "Stores main project details."],
        ["phases", "id, project_id, title, description, order", "Stores project phases."],
        ["activities", "id, phase_id, title, description", "Stores activities under phases."],
        ["tasks", "id, activity_id, assigned_to, title, task_type, status, estimated_hours, start_date, end_date, dependencies, skills_required, is_critical, slack", "Stores individual task details."],
        ["team_members", "id, name, role, skills, experience", "Stores team member information."],
        ["project_members", "project_id, team_member_id", "Junction table for project-team relationship."],
    ], [1.4, 3.1, 1.8])
    add_para(doc, "The design is mostly normalized because projects, phases, activities, tasks, and team members are stored in separate tables. However, some multi-value fields such as skills, dependencies, and required skills are currently stored as comma-separated text. Future versions can normalize these into separate relational tables.")

    doc.add_heading("Forms and Report Design", level=2)
    add_para(doc, "The future frontend can include forms for entering project metadata, uploading the scope document, and adding team members. Reports can display the generated WBS, scheduled tasks, task assignments, and critical path.")
    add_figure_placeholder(doc, "Figure 6: Form and Report Design for ScaffoldPM", "Create a combined UI wireframe. Left side: Generate Project form with project name, start date, end date, file upload, and team member entry fields. Right side: Generated Report view showing project summary, WBS tree, task table with assignee, estimated hours, start date, end date, status, and critical flag.")
    add_caption(doc, "Figure 6: Form and Report Design for ScaffoldPM")

    doc.add_heading("Interface and Dialogue Design", level=2)
    add_para(doc, "The interface dialogue begins when the user opens the project generation screen, uploads a scope document, fills project metadata, and submits the form. The backend returns the generated project plan, which can be displayed as a WBS tree, task board, schedule, and critical path summary.")
    add_figure_placeholder(doc, "Figure 7: Interface and Dialogue Design for ScaffoldPM", "Create a dialogue diagram with screens: Project List, Generate Project, Team Member Input, Upload Scope Document, Generated WBS View, Task Detail View, Kanban Status View, and Critical Path View. Show navigation arrows between screens and API interactions with the backend.")
    add_caption(doc, "Figure 7: Interface and Dialogue Design for ScaffoldPM")

    doc.add_heading("API Design", level=2)
    add_table(doc, ["Endpoint", "Method", "Purpose"], [
        ["/api/health", "GET", "Check backend health status."],
        ["/api/health/extract", "POST", "Extract text from uploaded document for testing."],
        ["/api/project/generate-wbs", "POST", "Generate WBS, assignments, schedule, and critical path."],
        ["/api/project/", "GET", "List saved projects."],
        ["/api/project/{project_id}", "GET", "Fetch full project details."],
        ["/api/project/task/status", "PATCH", "Update task status."],
        ["/api/project/task/assign", "PATCH", "Update task assignment."],
    ], [2.2, 0.9, 3.2])

    doc.add_heading("Algorithm Description", level=2)
    doc.add_heading("Purpose", level=3)
    add_para(doc, "The algorithms used in ScaffoldPM are designed to transform unstructured scope information into a structured and optimized project plan.")

    doc.add_heading("Working Principle", level=3)
    add_numbers(doc, [
        "The document extraction algorithm identifies the uploaded file type and extracts text using suitable libraries.",
        "The WBS generation algorithm sends project information and cleaned scope text to an LLM and parses the response into a structured schema.",
        "The dependency resolution algorithm maps dependency task titles returned by the LLM into internal task IDs.",
        "The task matching algorithm creates embeddings for task profiles and team member profiles, then assigns each task using similarity score, experience multiplier, and workload penalty.",
        "The duration estimation algorithm loads a trained XGBoost model when available. If unavailable, it uses a heuristic formula based on task type, experience, skills, and description length.",
        "The scheduling algorithm places dependency-free tasks first and schedules remaining tasks when dependencies and assigned members are available.",
        "The CPM algorithm creates a directed graph, computes early and late timing values, calculates slack, and marks tasks with zero slack as critical.",
    ])

    doc.add_heading("Pseudo-Code for Main Workflow", level=3)
    code = [
        "Input: scope_document, project_payload",
        "Validate project_payload",
        "scope_text = extract_document_text(scope_document)",
        "project = generate_wbs(scope_text, project_payload)",
        "project = assign_tasks_to_members(project)",
        "project = estimate_task_durations(project)",
        "project = schedule_project(project)",
        "project, critical_path = analyze_critical_path(project)",
        "save_project(project)",
        "Return project, total_duration_days, critical_path",
    ]
    for line in code:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(line)
        set_font(r, name="Consolas", size=10)


def add_references(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("REFERENCES", level=1)
    refs = [
        '[1] Project Management Institute, A Guide to the Project Management Body of Knowledge (PMBOK Guide), 7th ed. Project Management Institute, 2021.',
        '[2] Project Management Institute, Practice Standard for Work Breakdown Structures, 3rd ed. Project Management Institute, 2019.',
        '[3] R. T. Hans, "Work Breakdown Structure: A Tool for Software Project Scope Verification," arXiv:1308.2876, 2013.',
        '[4] J. E. Kelley Jr. and M. R. Walker, "Critical-path planning and scheduling," in Proc. Eastern Joint IRE-AIEE-ACM Computer Conference, 1959.',
        '[5] N. Reimers and I. Gurevych, "Sentence-BERT: Sentence embeddings using Siamese BERT-networks," arXiv:1908.10084, 2019.',
        '[6] T. Chen and C. Guestrin, "XGBoost: A scalable tree boosting system," in Proc. 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining, 2016, pp. 785-794.',
        '[7] B. W. Boehm, Software Engineering Economics. Prentice-Hall, 1981.',
        '[8] A. Trendowicz, J. Munch, and R. Jeffery, "State of the practice in software effort estimation: A survey and literature review," arXiv:1401.5878, 2014.',
        '[9] Y. Mahmood, N. Kama, A. Azmi, A. S. Khan, and M. Ali, "Software effort estimation accuracy prediction of machine learning techniques: A systematic performance evaluation," arXiv:2101.10658, 2021.',
        '[10] FastAPI, "FastAPI documentation." https://fastapi.tiangolo.com/',
        '[11] Pydantic, "Pydantic documentation: Models." https://docs.pydantic.dev/latest/concepts/models/',
        '[12] SQLAlchemy, "SQLAlchemy 2.0 documentation: Asynchronous I/O." https://docs.sqlalchemy.org/en/20/orm/extensions/asyncio.html',
        '[13] SQLite, "About SQLite." https://www.sqlite.org/about.html',
        '[14] LangChain, "LangChain overview." https://docs.langchain.com/oss/python/langchain/overview',
        '[15] OpenAI, "Text generation guide." https://platform.openai.com/docs/guides/text',
        '[16] OpenAI, "Structured model outputs." https://platform.openai.com/docs/guides/structured-outputs',
        '[17] Sentence Transformers, "SentenceTransformers documentation." https://www.sbert.net/',
        '[18] XGBoost Developers, "XGBoost documentation." https://xgboost.readthedocs.io/',
        '[19] NetworkX Developers, "topological_sort - NetworkX documentation." https://networkx.org/documentation/stable/reference/algorithms/generated/networkx.algorithms.dag.topological_sort.html',
    ]
    for ref in refs:
        add_para(doc, ref, justify=True)


def main() -> None:
    doc = Document()
    configure_page(doc)
    setup_styles(doc)
    add_front_matter(doc)
    add_chapter_1(doc)
    add_chapter_2(doc)
    add_chapter_3(doc)
    add_chapter_4(doc)
    add_references(doc)
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
